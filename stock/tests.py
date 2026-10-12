from decimal import Decimal
from django.test import TestCase, Client
from django.urls import reverse
from django.utils import timezone
from django.contrib.auth import get_user_model
from stock.forms import ProductForm
from stock.importer import import_products_from_rows, read_rows_from_workbook
from stock.models import (
    Category, Supplier, Product, StockAudit, StockAuditItem, StockCountSession,
    PurchaseOrder, StockMovement, AppSettings,
)

User = get_user_model()


class RestaurantInventoryTests(TestCase):
    def setUp(self):
        self.client = Client()

        # Crear usuario para autenticación
        self.user = User.objects.create_user(username='chef_mario', password='password123')
        self.client.login(username='chef_mario', password='password123')

        # Crear categoría y proveedor
        self.category = Category.objects.create(name='Carnes & Aves')
        self.supplier = Supplier.objects.create(
            name='Carnes Selectas San José',
            phone='+502 5555 1234',
            email='pedidos@carnes.com'
        )

        # Crear productos con ordinal / secuencia definida
        self.product1 = Product.objects.create(
            name='Filete de Res',
            sku='RES-01',
            sequence_number=10,
            category=self.category,
            supplier=self.supplier,
            unit='kg',
            current_stock=Decimal('5.00'),
            min_stock=Decimal('10.00'),
            ideal_stock=Decimal('15.00'),
            cost_price=Decimal('300.00'),
            is_active=True
        )

        self.product2 = Product.objects.create(
            name='Pechuga de Pollo',
            sku='POLLO-01',
            sequence_number=20,
            category=self.category,
            supplier=self.supplier,
            unit='kg',
            current_stock=Decimal('3.00'),
            min_stock=Decimal('15.00'),
            ideal_stock=Decimal('20.00'),
            cost_price=Decimal('120.00'),
            is_active=True
        )

    def test_product_ordering_by_sequence(self):
        """Verifica que los productos se ordenan por su ordinal/secuencia."""
        products = list(Product.objects.filter(is_active=True))
        self.assertEqual(products[0].sequence_number, 10)
        self.assertEqual(products[1].sequence_number, 20)

    def test_product_form_supports_multiple_sections_and_optional_fields(self):
        second_category = Category.objects.create(name='Barra')
        form = ProductForm(data={
            'name': 'Producto sin datos opcionales',
            'sequence_number': 30,
            'location': [self.category.pk, second_category.pk],
            'sku': '',
            'supplier': '',
            'unit': '',
            'current_stock': '',
            'min_stock': '',
            'ideal_stock': '',
            'cost_price': '',
            'detalle_ubicacion': '',
            'is_active': 'on',
        })

        self.assertTrue(form.is_valid(), form.errors)
        product = form.save()
        self.assertEqual(set(product.location.all()), {self.category, second_category})
        self.assertIsNone(product.current_stock)
        self.assertIsNone(product.min_stock)
        self.assertIsNone(product.ideal_stock)
        self.assertIsNone(product.cost_price)

    def test_product_form_renders_searchable_multi_section_selector(self):
        Category.objects.create(name='Barra')

        response = self.client.get(reverse('stock:product_create'))

        self.assertContains(response, 'id="categoryMultiSelectButton"')
        self.assertContains(response, 'id="categoryMultiSelectSearch"')
        self.assertContains(response, 'type="checkbox"')
        self.assertContains(response, 'Carnes &amp; Aves')
        self.assertContains(response, 'Barra')

    def test_null_optional_product_fields_do_not_break_pages_or_mark_stock_zero(self):
        product = Product.objects.create(
            name='Producto con información pendiente',
            sequence_number=30,
            current_stock=None,
            min_stock=None,
            ideal_stock=None,
            cost_price=None,
        )

        self.assertFalse(product.is_out_of_stock)
        self.assertFalse(product.is_below_minimum)
        self.assertEqual(product.suggested_reorder_qty, Decimal('0.00'))
        self.assertEqual(self.client.get(reverse('stock:dashboard')).status_code, 200)
        self.assertEqual(self.client.get(reverse('stock:product_list')).status_code, 200)
        self.assertEqual(self.client.get(reverse('stock:product_detail', kwargs={'pk': product.pk})).status_code, 200)

    def test_dashboard_puts_count_sessions_before_critical_stock_and_paginates(self):
        for index in range(9):
            StockCountSession.objects.create(
                title=f'Registro {index}', count_date=timezone.now().date(),
            )

        page = self.client.get(reverse('stock:dashboard'))
        body = page.content.decode()
        self.assertLess(body.index('Conteos de Stock Recientes'), body.index('Insumos Críticos'))
        self.assertContains(page, 'Registro 8')
        self.assertNotContains(page, 'Registro 0')
        self.assertContains(page, 'Siguiente')

        page_two = self.client.get(reverse('stock:dashboard') + '?conteos=2')
        self.assertContains(page_two, 'Registro 0')
        self.assertNotContains(page_two, 'Registro 8')
        self.assertContains(page_two, 'Anterior')

    def test_product_cost_is_displayed_in_quetzales(self):
        response = self.client.get(reverse('stock:product_list'))
        self.assertContains(response, 'Q300.00')
        self.assertNotContains(response, '$300.00')

    def test_product_can_be_found_through_any_assigned_section(self):
        second_category = Category.objects.create(name='Barra')
        self.product1.location.add(second_category)

        response = self.client.get(
            reverse('stock:product_list'),
            {'category': second_category.pk},
        )

        self.assertContains(response, 'Filete de Res')
        self.assertNotContains(response, 'Pechuga de Pollo')

    def test_audit_section_filter_includes_products_assigned_to_multiple_sections(self):
        second_category = Category.objects.create(name='Barra')
        self.product1.location.add(second_category)

        response = self.client.post(reverse('stock:stock_audit_create'), {
            'title': 'Conteo de barra',
            'audit_date': timezone.now().date().isoformat(),
            'category': second_category.pk,
            'location_area': 'Barra',
            'responsible': '',
            'notes': '',
        })

        self.assertEqual(response.status_code, 302)
        audit = StockAudit.objects.get(title='Conteo de barra')
        self.assertQuerySetEqual(
            audit.items.order_by('product_id').values_list('product_id', flat=True),
            [self.product1.pk],
        )

    def test_order_report_handles_products_without_cost(self):
        product = Product.objects.create(
            name='Costo pendiente',
            sequence_number=30,
            current_stock=None,
            min_stock=None,
            ideal_stock=None,
            cost_price=None,
        )
        audit = StockAudit.objects.create(
            title='Conteo costo pendiente',
            audit_date=timezone.now().date(),
        )
        StockAuditItem.objects.create(
            audit=audit,
            product=product,
            final_order=Decimal('2.00'),
        )

        response = self.client.get(
            reverse('stock:stock_audit_order_report', kwargs={'pk': audit.pk})
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Q0.00')

    def test_unlocalized_count_values_in_input(self):
        """Verifica que los conteos registrados se muestran con punto en las casillas <input type='number'>."""
        audit = StockAudit.objects.create(
            title='Conteo Bodega',
            audit_date=timezone.now().date(),
            location_area='Bodega',
            status='draft'
        )
        item = StockAuditItem.objects.create(
            audit=audit,
            product=self.product2,
            previous_stock=Decimal('3.00'),
            min_stock=Decimal('15.00'),
            ideal_stock=Decimal('20.00'),
            counted_stock=Decimal('4.50'),
            suggested_order=Decimal('15.50'),
            final_order=Decimal('16.00')
        )

        resp = self.client.get(reverse('stock:stock_audit_count', kwargs={'pk': audit.pk}))
        self.assertEqual(resp.status_code, 200)
        # Debe contener value="4.50" (con punto, gracias a unlocalize) para que el navegador lo muestre en input number
        self.assertContains(resp, 'value="4.50"')
        self.assertContains(resp, 'value="16.00"')

    def test_purchase_order_lock_and_unlock(self):
        """Verifica la funcionalidad de bloquear y desbloquear pedidos ya hechos."""
        po = PurchaseOrder.objects.create(
            code='PED-TEST-001',
            supplier=self.supplier,
            date=timezone.now().date(),
            status='pending'
        )

        # 1. Bloquear pedido
        lock_url = reverse('stock:purchase_order_toggle_lock', kwargs={'pk': po.pk})
        self.client.post(lock_url)
        po.refresh_from_db()
        self.assertTrue(po.is_locked)
        self.assertEqual(po.locked_by, 'chef_mario')

        # 2. Intentar cambiar estado mientras está bloqueado (debe ser rechazado)
        status_url = reverse('stock:purchase_order_change_status', kwargs={'pk': po.pk})
        self.client.post(status_url, {'status': 'sent'})
        po.refresh_from_db()
        self.assertEqual(po.status, 'pending') # Sigue en pending

        # 3. Desbloquear pedido
        self.client.post(lock_url)
        po.refresh_from_db()
        self.assertFalse(po.is_locked)

        # 4. Ahora sí permite cambiar estado
        self.client.post(status_url, {'status': 'sent'})
        po.refresh_from_db()
        self.assertEqual(po.status, 'sent')

    def test_consolidated_daily_report(self):
        """
        Verifica la consolidación de conteos de múltiples ubicaciones
        (ej: Bodega y Barra para el mismo producto).
        """
        # Toma 1: Bodega Central
        audit_bodega = StockAudit.objects.create(
            title='Conteo Diario Bodega',
            audit_date=timezone.now().date(),
            location_area='Bodega Central',
            status='draft'
        )
        StockAuditItem.objects.create(
            audit=audit_bodega,
            product=self.product1,
            counted_stock=Decimal('4.00')  # 4 kg en bodega
        )

        # Toma 2: Cocina Principal
        audit_cocina = StockAudit.objects.create(
            title='Conteo Diario Cocina',
            audit_date=timezone.now().date(),
            location_area='Cocina',
            status='draft'
        )
        StockAuditItem.objects.create(
            audit=audit_cocina,
            product=self.product1,
            counted_stock=Decimal('2.00')  # 2 kg en cocina
        )

        # Consolidar ambas tomas
        resp = self.client.post(reverse('stock:consolidated_report_view'), {
            'audit_ids': [audit_bodega.id, audit_cocina.id]
        })
        self.assertEqual(resp.status_code, 200)

        # Total contado: 4.00 + 2.00 = 6.00 kg.
        # Min stock es 10.00, ideal es 15.00 -> Cantidad a pedir sugerida = 15.00 - 6.00 = 9.00 kg
        self.assertContains(resp, 'Filete de Res')
        self.assertContains(resp, 'Bodega Central')
        self.assertContains(resp, 'Cocina')


    def test_session_create_builds_one_audit_per_area(self):
        """Crear un registro de conteo genera un conteo por cada área seleccionada."""
        barra = Category.objects.create(name='Barra')
        self.product1.location.add(barra)

        response = self.client.post(reverse('stock:stock_count_session_create'), {
            'title': 'Conteo General Test',
            'count_date': timezone.now().date().isoformat(),
            'responsible': 'Chef Mario',
            'notes': '',
            'areas': [self.category.pk, barra.pk],
        })

        self.assertEqual(response.status_code, 302)
        session = StockCountSession.objects.get(title='Conteo General Test')
        self.assertEqual(session.audits.count(), 2)
        self.assertEqual(
            set(session.audits.values_list('location_area', flat=True)),
            {'Carnes & Aves', 'Barra'},
        )
        carnes_audit = session.audits.get(location_area='Carnes & Aves')
        self.assertEqual(
            set(carnes_audit.items.values_list('product_id', flat=True)),
            {self.product1.pk, self.product2.pk},
        )

    def test_session_finalize_consolidates_multiple_areas_into_inventory(self):
        """El stock solo se aplica al finalizar, sumando las áreas del mismo producto."""
        barra = Category.objects.create(name='Barra')
        self.product1.location.add(barra)  # product1 vive en dos áreas

        session = StockCountSession.objects.create(
            title='Conteo Multi Área',
            count_date=timezone.now().date(),
        )
        audit_a = StockAudit.objects.create(
            session=session, category=self.category, location_area='Carnes & Aves',
            title='Conteo Multi Área - Carnes', audit_date=session.count_date,
        )
        audit_b = StockAudit.objects.create(
            session=session, category=barra, location_area='Barra',
            title='Conteo Multi Área - Barra', audit_date=session.count_date,
        )
        StockAuditItem.objects.create(audit=audit_a, product=self.product1, counted_stock=Decimal('4.00'))
        StockAuditItem.objects.create(audit=audit_b, product=self.product1, counted_stock=Decimal('2.00'))

        # El stock no cambia hasta finalizar la sesión
        self.assertEqual(Product.objects.get(pk=self.product1.pk).current_stock, Decimal('5.00'))

        resp = self.client.post(reverse('stock:stock_count_session_finalize', kwargs={'pk': session.pk}))
        self.assertEqual(resp.status_code, 302)

        session.refresh_from_db()
        self.product1.refresh_from_db()
        self.assertEqual(session.status, 'completed')
        self.assertEqual(self.product1.current_stock, Decimal('6.00'))
        self.assertTrue(
            StockMovement.objects.filter(
                product=self.product1, reference__icontains='Consolidado'
            ).exists()
        )

    def test_session_area_complete_does_not_apply_stock_immediately(self):
        """Finalizar un área de una sesión no toca el inventario hasta finalizar la sesión."""
        session = StockCountSession.objects.create(
            title='Conteo por Área',
            count_date=timezone.now().date(),
        )
        audit = StockAudit.objects.create(
            session=session, category=self.category, location_area='Carnes & Aves',
            title='Conteo por Área - Carnes', audit_date=session.count_date,
        )
        item = StockAuditItem.objects.create(
            audit=audit, product=self.product1,
            min_stock=Decimal('10.00'), ideal_stock=Decimal('15.00'),
        )

        resp = self.client.post(reverse('stock:stock_audit_count', kwargs={'pk': audit.pk}), {
            f'counted_{item.pk}': '7',
            'action': 'complete',
        })

        self.assertEqual(resp.status_code, 302)
        audit.refresh_from_db()
        self.product1.refresh_from_db()
        self.assertEqual(audit.status, 'completed')
        self.assertEqual(self.product1.current_stock, Decimal('5.00'))

    def test_session_detail_and_list_render(self):
        session = StockCountSession.objects.create(
            title='Conteo Visible', count_date=timezone.now().date(),
        )
        StockAudit.objects.create(
            session=session, category=self.category, location_area='Carnes & Aves',
            title='Conteo Visible - Carnes', audit_date=session.count_date,
        )

        self.assertEqual(self.client.get(reverse('stock:stock_count_session_list')).status_code, 200)
        detail = self.client.get(
            reverse('stock:stock_count_session_detail', kwargs={'pk': session.pk})
        )
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, 'Carnes &amp; Aves')
        self.assertContains(detail, 'Finalizar y Consolidar')

    def test_unauthenticated_user_redirected_to_login(self):
        """Verifica que si no está autenticado, redirige al login."""
        self.client.logout()
        resp = self.client.get(reverse('stock:dashboard'))
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/accounts/login/', resp.url)

    def test_import_products_from_rows_creates_areas_and_units(self):
        rows = [
            {'sequence_number': 5, 'group': 'Abarrotes', 'name': 'Arroz Blanco',
             'presentation': 'Saco 50 lb', 'location': 'Bodega, Cocina'},
            {'name': 'Aceite Vegetal', 'presentation': 'Galón'},
        ]
        summary = import_products_from_rows(rows)

        self.assertEqual(summary['created'], 2)
        self.assertEqual(summary['errors'], [])

        arroz = Product.objects.get(name='Arroz Blanco')
        self.assertEqual(arroz.unit, 'unidad')
        self.assertEqual(arroz.group, 'Abarrotes')
        self.assertEqual(arroz.presentation, 'Saco 50 lb')
        self.assertEqual(arroz.sequence_number, 5)
        self.assertEqual(
            set(arroz.location.values_list('name', flat=True)), {'Bodega', 'Cocina'}
        )

        aceite = Product.objects.get(name='Aceite Vegetal')
        self.assertEqual(aceite.unit, 'unidad')
        self.assertEqual(aceite.location.count(), 0)

    def test_import_products_from_rows_updates_existing_by_name(self):
        rows = [{'name': 'filete de res', 'group': 'Carnes', 'presentation': 'Libra'}]
        summary = import_products_from_rows(rows)

        self.assertEqual(summary['updated'], 1)
        self.product1.refresh_from_db()
        self.assertEqual(self.product1.group, 'Carnes')
        self.assertEqual(self.product1.presentation, 'Libra')
        self.assertEqual(Product.objects.filter(name__iexact='filete de res').count(), 1)

    def test_count_lists_only_products_of_the_area(self):
        """El conteo de un área solo muestra los productos asignados a esa área."""
        barra = Category.objects.create(name='Barra')
        cerveza = Product.objects.create(name='Cerveza Artesanal', sequence_number=40, unit='unidad')
        cerveza.location.add(barra)

        session = StockCountSession.objects.create(
            title='Conteo Barra', count_date=timezone.now().date()
        )
        audit = StockAudit.objects.create(
            session=session, category=barra, location_area='Barra',
            title='Conteo Barra', audit_date=timezone.now().date(),
        )

        response = self.client.get(reverse('stock:stock_audit_count', kwargs={'pk': audit.pk}))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Cerveza Artesanal')
        self.assertNotContains(response, 'Filete de Res')
        self.assertNotContains(response, 'Pechuga de Pollo')

    def test_read_rows_from_workbook_maps_headers(self):
        from openpyxl import Workbook
        import io as _io

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['sequence_number', 'group', 'name', 'presentation', 'location'])
        sheet.append([1, 'Bebidas', 'Coca Cola', 'Botella 3L', 'Barra'])
        buffer = _io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        rows = read_rows_from_workbook(buffer)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['name'], 'Coca Cola')
        self.assertEqual(rows[0]['group'], 'Bebidas')
        self.assertEqual(rows[0]['location'], 'Barra')

    def test_import_reads_sku_column(self):
        from openpyxl import Workbook
        import io as _io

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['sku', 'group', 'name', 'presentation', 'location'])
        sheet.append(['ABC-123', 'Carnes', 'Solomillo', 'Libra', 'Cocina'])
        buffer = _io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        rows = read_rows_from_workbook(buffer)
        self.assertEqual(rows[0].get('sku'), 'ABC-123')

        summary = import_products_from_rows(rows)
        self.assertEqual(summary['created'], 1)
        prod = Product.objects.get(name='Solomillo')
        self.assertEqual(prod.sku, 'ABC-123')

    def test_count_marked_checked_without_quantity_sets_zero(self):
        """Marcar el checkbox 'Contado' sin indicar cantidad establece counted_stock en 0."""
        session = StockCountSession.objects.create(
            title='Conteo Checkbox', count_date=timezone.now().date()
        )
        audit = StockAudit.objects.create(
            session=session, category=self.category, location_area='Carnes & Aves',
            title='Conteo Checkbox - Carnes', audit_date=session.count_date,
        )
        item = StockAuditItem.objects.create(
            audit=audit, product=self.product1,
            min_stock=Decimal('10.00'), ideal_stock=Decimal('15.00'),
        )

        resp = self.client.post(reverse('stock:stock_audit_count', kwargs={'pk': audit.pk}), {
            f'checked_{item.pk}': 'on',
        })

        self.assertEqual(resp.status_code, 302)
        item.refresh_from_db()
        self.assertEqual(item.counted_stock, Decimal('0.00'))
        self.assertEqual(item.suggested_order, Decimal('15.00'))

    def test_count_unchecked_remains_pending(self):
        """Dejar sin marcar mantiene counted_stock como None."""
        session = StockCountSession.objects.create(
            title='Conteo Pendiente', count_date=timezone.now().date()
        )
        audit = StockAudit.objects.create(
            session=session, category=self.category, location_area='Carnes & Aves',
            title='Conteo Pendiente - Carnes', audit_date=session.count_date,
        )
        item = StockAuditItem.objects.create(
            audit=audit, product=self.product1,
            min_stock=Decimal('10.00'), ideal_stock=Decimal('15.00'),
            counted_stock=Decimal('5.00'),
        )

        resp = self.client.post(reverse('stock:stock_audit_count', kwargs={'pk': audit.pk}), {
            # No enviamos ni cantidad ni checked
        })

        self.assertEqual(resp.status_code, 302)
        item.refresh_from_db()
        self.assertIsNone(item.counted_stock)

    def test_count_fields_show_labels_above_fillable_inputs(self):
        product = Product.objects.create(
            name='Hielo', sequence_number=50, unit='unidad', min_stock=None,
        )
        audit = StockAudit.objects.create(
            title='Conteo etiquetas', audit_date=timezone.now().date(), status='draft',
        )
        StockAuditItem.objects.create(
            audit=audit, product=product, previous_stock=Decimal('2.00'),
        )

        response = self.client.get(reverse('stock:stock_audit_count', kwargs={'pk': audit.pk}))

        self.assertContains(response, 'Cantidad a contar')
        self.assertContains(response, 'Cantidad a pedir')
        self.assertContains(response, 'Escribe la cantidad')
        self.assertContains(response, 'Sin mínimo')
        self.assertContains(response, 'Esta acción colocará el conteo a cero de')
        self.assertNotContains(response, 'placeholder="—"')

    def test_saving_one_product_keeps_progress_for_later(self):
        audit = StockAudit.objects.create(
            title='Conteo progreso', audit_date=timezone.now().date(), status='draft',
        )
        item = StockAuditItem.objects.create(
            audit=audit, product=self.product1,
            previous_stock=Decimal('5.00'), min_stock=Decimal('10.00'), ideal_stock=Decimal('15.00'),
        )
        other = StockAuditItem.objects.create(
            audit=audit, product=self.product2,
            previous_stock=Decimal('3.00'), min_stock=Decimal('15.00'), ideal_stock=Decimal('20.00'),
        )

        response = self.client.post(
            reverse('stock:stock_audit_item_save', kwargs={'pk': audit.pk, 'item_pk': item.pk}),
            {'counted': '7.25', 'order': '3.50', 'notes': 'cámara 1'},
        )

        self.assertEqual(response.status_code, 200)
        item.refresh_from_db()
        other.refresh_from_db()
        self.assertEqual(item.counted_stock, Decimal('7.25'))
        self.assertEqual(item.final_order, Decimal('3.50'))
        self.assertEqual(item.notes, 'cámara 1')
        self.assertIsNone(other.counted_stock)

        page = self.client.get(reverse('stock:stock_audit_count', kwargs={'pk': audit.pk}))
        self.assertContains(page, 'value="7.25"')
        self.assertContains(page, 'cámara 1')

    def test_mark_pending_zero_updates_only_selected_products(self):
        audit = StockAudit.objects.create(
            title='Conteo ceros', audit_date=timezone.now().date(), status='draft',
        )
        pending = StockAuditItem.objects.create(
            audit=audit, product=self.product1,
            min_stock=Decimal('10.00'), ideal_stock=Decimal('15.00'),
        )
        already = StockAuditItem.objects.create(
            audit=audit, product=self.product2,
            min_stock=Decimal('15.00'), ideal_stock=Decimal('20.00'),
            counted_stock=Decimal('4.00'), final_order=Decimal('2.00'),
        )

        response = self.client.post(
            reverse('stock:stock_audit_mark_pending_zero', kwargs={'pk': audit.pk}),
            {'item_ids': [pending.pk]},
        )

        self.assertEqual(response.status_code, 200)
        pending.refresh_from_db()
        already.refresh_from_db()
        self.assertEqual(pending.counted_stock, Decimal('0.00'))
        self.assertEqual(pending.suggested_order, Decimal('15.00'))
        self.assertEqual(already.counted_stock, Decimal('4.00'))
        self.assertEqual(already.final_order, Decimal('2.00'))

    def test_session_excel_puts_each_area_in_its_own_column(self):
        from io import BytesIO
        from openpyxl import load_workbook

        bodega = Category.objects.create(name='Bodega')
        cocina = Category.objects.create(name='Cocina')
        self.product1.group = 'ABARROTE'
        self.product1.presentation = 'bolsa'
        self.product1.save(update_fields=['group', 'presentation'])

        session = StockCountSession.objects.create(
            title='SUSHI MON', count_date=timezone.now().date(),
        )
        audit_bodega = StockAudit.objects.create(
            session=session, category=bodega, location_area='Bodega',
            title='SUSHI MON - Bodega', audit_date=session.count_date,
        )
        audit_cocina = StockAudit.objects.create(
            session=session, category=cocina, location_area='Cocina',
            title='SUSHI MON - Cocina', audit_date=session.count_date,
        )
        StockAuditItem.objects.create(
            audit=audit_bodega, product=self.product2, counted_stock=Decimal('1.00'),
        )
        StockAuditItem.objects.create(
            audit=audit_bodega, product=self.product1, counted_stock=None,
            final_order=Decimal('0.75'),
        )
        StockAuditItem.objects.create(
            audit=audit_cocina, product=self.product1, counted_stock=Decimal('0.50'),
            final_order=Decimal('1.25'),
        )

        detail = self.client.get(reverse('stock:stock_count_session_detail', kwargs={'pk': session.pk}))
        self.assertContains(detail, 'Descargar Excel')
        self.assertContains(detail, 'Bodega')
        self.assertContains(detail, 'Cocina')

        response = self.client.post(
            reverse('stock:stock_count_session_export_excel', kwargs={'pk': session.pk}),
            {'areas': [audit_cocina.pk, audit_bodega.pk]},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])

        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet['A1'].value, 'SUSHI MON')
        headers = [cell.value for cell in sheet[2]]
        self.assertEqual(headers, [
            'Orden', 'Grupo', 'Código', 'Nombre del producto', 'Presentación',
            'Bodega', 'Cocina', 'Total existencias', 'Pedido',
        ])
        self.assertEqual([sheet.cell(3, col).value for col in range(1, 6)], [
            10, 'ABARROTE', 'RES-01', 'Filete de Res', 'bolsa',
        ])
        self.assertEqual(sheet.cell(3, 6).value, Decimal('0.00'))
        self.assertEqual(sheet.cell(3, 6).number_format, '"0.00"')
        self.assertEqual(sheet.cell(3, 7).number_format, '"0.50"')
        self.assertEqual(sheet.cell(3, 8).number_format, '"0.50"')
        self.assertEqual(sheet.cell(3, 7).value, Decimal('0.50'))
        self.assertEqual(sheet.cell(3, 8).value, '=SUM(F3:G3)')
        self.assertEqual(sheet.cell(3, 9).value, Decimal('2.00'))
        self.assertEqual(sheet.cell(3, 9).number_format, '"2.00"')
        self.assertEqual(sheet.cell(4, 4).value, 'Pechuga de Pollo')
        self.assertEqual(sheet.cell(4, 6).value, Decimal('1.00'))
        self.assertEqual(sheet.cell(4, 7).value, Decimal('0.00'))
        self.assertIsNone(sheet.cell(4, 9).value)

        only_cocina = self.client.post(
            reverse('stock:stock_count_session_export_excel', kwargs={'pk': session.pk}),
            {'areas': [audit_cocina.pk]},
        )
        cocina_sheet = load_workbook(BytesIO(only_cocina.content)).active
        cocina_headers = [cell.value for cell in cocina_sheet[2]]
        self.assertIn('Cocina', cocina_headers)
        self.assertNotIn('Bodega', cocina_headers)
        self.assertEqual(cocina_headers[-1], 'Pedido')
        self.assertEqual(cocina_sheet.cell(3, 8).value, Decimal('1.25'))

        missing = self.client.post(
            reverse('stock:stock_count_session_export_excel', kwargs={'pk': session.pk}),
            {},
        )
        self.assertEqual(missing.status_code, 302)

        comma_file = self.client.post(
            reverse('stock:stock_count_session_export_excel', kwargs={'pk': session.pk}),
            {'areas': [audit_cocina.pk], 'decimal_separator': ','},
        )
        comma_sheet = load_workbook(BytesIO(comma_file.content)).active
        self.assertEqual(comma_sheet.cell(3, 6).number_format, '"0,50"')
        self.assertEqual(comma_sheet.cell(3, 6).value, Decimal('0.50'))

    def test_settings_hide_previous_and_minimum_and_skip_order(self):
        AppSettings.objects.create(pk=1, hide_previous_field=True, hide_minimum_field=True)
        audit = StockAudit.objects.create(
            title='Conteo sin mínimo', audit_date=timezone.now().date(), status='draft',
        )
        item = StockAuditItem.objects.create(
            audit=audit, product=self.product1,
            previous_stock=Decimal('5.00'), min_stock=Decimal('10.00'), ideal_stock=Decimal('15.00'),
        )

        page = self.client.get(reverse('stock:stock_audit_count', kwargs={'pk': audit.pk}))
        self.assertContains(page, 'data-calculate-order="0"')
        self.assertNotContains(page, 'count-label">Mínimo')
        self.assertNotContains(page, 'count-label">Previo')
        self.assertNotContains(page, 'Sugerido')

        saved = self.client.post(
            reverse('stock:stock_audit_item_save', kwargs={'pk': audit.pk, 'item_pk': item.pk}),
            {'counted': '4', 'order': '', 'notes': ''},
        )
        self.assertEqual(saved.status_code, 200)
        item.refresh_from_db()
        self.assertEqual(item.counted_stock, Decimal('4.00'))
        self.assertEqual(item.suggested_order, Decimal('0.00'))
        self.assertEqual(item.final_order, Decimal('0.00'))

        typed = self.client.post(
            reverse('stock:stock_audit_item_save', kwargs={'pk': audit.pk, 'item_pk': item.pk}),
            {'counted': '4', 'order': '6', 'notes': ''},
        )
        self.assertEqual(typed.status_code, 200)
        item.refresh_from_db()
        self.assertEqual(item.final_order, Decimal('6.00'))

    def test_settings_page_saves_the_two_options(self):
        response = self.client.post(reverse('stock:app_settings'), {
            'hide_previous_field': 'on',
            'hide_minimum_field': 'on',
        })
        self.assertEqual(response.status_code, 302)
        settings_row = AppSettings.load()
        self.assertTrue(settings_row.hide_previous_field)
        self.assertTrue(settings_row.hide_minimum_field)

        cleared = self.client.post(reverse('stock:app_settings'), {})
        self.assertEqual(cleared.status_code, 302)
        settings_row.refresh_from_db()
        self.assertFalse(settings_row.hide_previous_field)
        self.assertFalse(settings_row.hide_minimum_field)

    def test_count_shows_chip_when_another_area_already_counted(self):
        session = StockCountSession.objects.create(
            title='Conteo compartido', count_date=timezone.now().date(),
        )
        bodega = StockAudit.objects.create(
            session=session, location_area='Bodega', title='Bodega',
            audit_date=session.count_date,
        )
        cocina = StockAudit.objects.create(
            session=session, location_area='Cocina', title='Cocina',
            audit_date=session.count_date,
        )
        StockAuditItem.objects.create(
            audit=bodega, product=self.product1, counted_stock=Decimal('10.00'),
        )
        StockAuditItem.objects.create(
            audit=cocina, product=self.product1, counted_stock=None,
        )
        StockAuditItem.objects.create(
            audit=bodega, product=self.product2, counted_stock=None,
        )
        StockAuditItem.objects.create(
            audit=cocina, product=self.product2, counted_stock=None,
        )

        cocina_page = self.client.get(reverse('stock:stock_audit_count', kwargs={'pk': cocina.pk}))
        self.assertContains(cocina_page, 'El área "Bodega" ha contado:')
        bodega_page = self.client.get(reverse('stock:stock_audit_count', kwargs={'pk': bodega.pk}))
        self.assertNotContains(bodega_page, 'El área "Bodega" ha contado:')
