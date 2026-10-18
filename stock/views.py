import csv
from decimal import Decimal
from django.core.paginator import Paginator
from django.db.models import prefetch_related_objects
from django.shortcuts import render, get_object_or_404, redirect
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from django.db.models import Q, Sum, F, Count
from django.utils import timezone
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_POST
from .models import (
    Product, Category, Supplier, StockAudit, StockAuditItem, StockCountSession,
    PurchaseOrder, PurchaseOrderItem, StockMovement, AppSettings, ProductAreaSequence,
    order_queryset_by_area, sequences_for_area,
)
from .forms import (
    ProductForm, StockAuditForm, StockCountSessionForm, SupplierForm, CategoryForm,
    QuickAdjustmentForm, AppSettingsForm,
)
from .importer import import_products_from_workbook


def _populate_audit_items(audit, category=None):
    """
    Precarga los items de conteo de una toma con los productos activos
    (opcionalmente filtrados por la categoría/área indicada) y retorna el conteo.
    """
    products_qs = Product.objects.filter(is_active=True)
    if category:
        products_qs = products_qs.filter(
            Q(location=category) | Q(category=category)
        ).distinct()
    products_qs = order_queryset_by_area(products_qs, category, product_id_field='pk', name_field='name')

    items_to_create = [
        StockAuditItem(
            audit=audit,
            product=prod,
            previous_stock=prod.current_stock or Decimal('0.00'),
            min_stock=prod.min_stock or Decimal('0.00'),
            ideal_stock=prod.ideal_stock or Decimal('0.00'),
            counted_stock=None,
            suggested_order=Decimal('0.00'),
            final_order=Decimal('0.00'),
        )
        for prod in products_qs
    ]
    StockAuditItem.objects.bulk_create(items_to_create)
    return len(items_to_create)


def _sync_area_items(audit, area):
    """
    Garantiza que el conteo (audit) tenga una fila por cada producto activo
    asignado al área indicada, de modo que solo se listen los productos del área.
    """
    if area is None:
        return 0
    existing_product_ids = set(audit.items.values_list('product_id', flat=True))
    products_qs = order_queryset_by_area(
        Product.objects.filter(is_active=True).filter(
            Q(location=area) | Q(category=area)
        ).distinct(),
        area,
        product_id_field='pk',
        name_field='name',
    )

    items_to_create = [
        StockAuditItem(
            audit=audit,
            product=prod,
            previous_stock=prod.current_stock or Decimal('0.00'),
            min_stock=prod.min_stock or Decimal('0.00'),
            ideal_stock=prod.ideal_stock or Decimal('0.00'),
            counted_stock=None,
            suggested_order=Decimal('0.00'),
            final_order=Decimal('0.00'),
        )
        for prod in products_qs
        if prod.id not in existing_product_ids
    ]
    if items_to_create:
        StockAuditItem.objects.bulk_create(items_to_create)
    return len(items_to_create)


def _orders_are_calculated():
    """El pedido se calcula solo cuando el campo Mínimo está visible."""
    return not AppSettings.load().hide_minimum_field


def _other_area_counts(audit, product_ids):
    """Conteos ya hechos del mismo producto en otras áreas de este conteo general."""
    if not audit.session_id or not product_ids:
        return {}
    rows = (
        StockAuditItem.objects.filter(
            audit__session_id=audit.session_id,
            product_id__in=product_ids,
            counted_stock__isnull=False,
        )
        .exclude(audit_id=audit.id)
        .select_related('audit', 'audit__category')
        .order_by('audit__location_area', 'audit_id')
    )
    grouped = {}
    for row in rows:
        area_name = row.audit.location_area or (
            row.audit.category.name if row.audit.category_id else row.audit.title
        )
        grouped.setdefault(row.product_id, []).append({
            'area': area_name,
            'counted': row.counted_stock,
        })
    return grouped


def _apply_item_capture(item, counted_raw, order_raw, notes_raw, is_checked=False, calculate_order=None):
    """Guarda el conteo, el pedido y las notas de un producto."""
    if calculate_order is None:
        calculate_order = _orders_are_calculated()
    target = item.ideal_stock if item.ideal_stock > 0 else item.min_stock
    counted_val = '' if counted_raw is None else str(counted_raw).strip().replace(',', '.')
    order_val = '' if order_raw is None else str(order_raw).strip().replace(',', '.')
    notes_val = '' if notes_raw is None else str(notes_raw).strip()

    parsed_order = None
    if order_val != '':
        try:
            parsed_order = max(Decimal('0.00'), Decimal(order_val))
        except Exception:
            parsed_order = None

    if counted_val != '':
        try:
            item.counted_stock = max(Decimal('0.00'), Decimal(counted_val))
        except Exception:
            item.counted_stock = Decimal('0.00')
        if calculate_order:
            suggested = max(Decimal('0.00'), target - item.counted_stock)
            item.suggested_order = suggested
            item.final_order = parsed_order if parsed_order is not None else suggested
        else:
            item.suggested_order = Decimal('0.00')
            item.final_order = parsed_order if parsed_order is not None else Decimal('0.00')
    elif is_checked:
        item.counted_stock = Decimal('0.00')
        if calculate_order:
            item.suggested_order = max(Decimal('0.00'), target)
            item.final_order = parsed_order if parsed_order is not None else item.suggested_order
        else:
            item.suggested_order = Decimal('0.00')
            item.final_order = parsed_order if parsed_order is not None else Decimal('0.00')
    else:
        item.counted_stock = None
        item.suggested_order = Decimal('0.00')
        item.final_order = parsed_order if parsed_order is not None else Decimal('0.00')

    item.notes = notes_val
    item.save()
    return item


def _item_capture_payload(item):
    product_min = item.product.min_stock
    has_min = bool(product_min)
    below_min = (
        item.counted_stock is not None
        and has_min
        and item.counted_stock <= product_min
    )
    return {
        'id': item.id,
        'counted_stock': None if item.counted_stock is None else f'{item.counted_stock:.2f}',
        'final_order': f'{item.final_order:.2f}',
        'suggested_order': f'{item.suggested_order:.2f}',
        'pending': item.counted_stock is None,
        'below_min': below_min,
        'has_min': has_min,
    }


@login_required
def app_settings(request):
    """Opciones generales: qué campos se ven al contar y si el pedido se calcula."""
    settings_row = AppSettings.load()
    if request.method == 'POST':
        form = AppSettingsForm(request.POST, instance=settings_row)
        if form.is_valid():
            form.save()
            messages.success(request, 'Configuración guardada.')
            return redirect('stock:app_settings')
    else:
        form = AppSettingsForm(instance=settings_row)
    return render(request, 'stock/app_settings.html', {'form': form, 'title': 'Configuración'})


@login_required
def dashboard(request):
    """Vista principal con métricas clave para el restaurante."""
    active_products = Product.objects.filter(is_active=True).select_related(
        'category', 'supplier'
    ).prefetch_related('location')
    total_products = active_products.count()
    
    # Productos por debajo del mínimo
    low_stock_products = [p for p in active_products if p.is_below_minimum]
    out_of_stock_products = [p for p in active_products if p.is_out_of_stock]
    
    # Valor estimado total del inventario
    total_inventory_value = sum(
        (p.current_stock or Decimal('0.00')) * (p.cost_price or Decimal('0.00'))
        for p in active_products
    )
    
    # Conteos de stock: una página, con el total de áreas en la misma consulta.
    sessions = StockCountSession.objects.annotate(
        area_total=Count('audits', distinct=True),
        area_done=Count('audits', filter=Q(audits__status='completed'), distinct=True),
    ).order_by('-count_date', '-created_at')
    session_page = Paginator(sessions, 8).get_page(request.GET.get('conteos'))
    session_rows = list(session_page.object_list)
    prefetch_related_objects(session_rows, 'audits')
    session_page.object_list = session_rows
    # Tomas de inventario recientes
    recent_audits = StockAudit.objects.select_related('category')[:5]
    
    # Órdenes de compra recientes
    recent_orders = PurchaseOrder.objects.select_related('supplier')[:5]
    
    # Categorías para acceso rápido
    categories = Category.objects.all()

    context = {
        'total_products': total_products,
        'low_stock_count': len(low_stock_products),
        'out_of_stock_count': len(out_of_stock_products),
        'total_inventory_value': total_inventory_value,
        'low_stock_products': low_stock_products[:8],  # Mostrar los primeros 8 críticos
        'session_page': session_page,
        'recent_audits': recent_audits,
        'recent_orders': recent_orders,
        'categories': categories,
    }
    return render(request, 'stock/dashboard.html', context)


# ==========================================
# GESTIÓN DE PRODUCTOS
# ==========================================

@login_required
def product_list(request):
    """Listado de productos con filtros, buscador y ordenados por secuencia/ordinal."""
    query = request.GET.get('q', '').strip()
    category_id = request.GET.get('category', '')
    supplier_id = request.GET.get('supplier', '')
    status_filter = request.GET.get('status', '')

    products = Product.objects.filter(is_active=True).select_related(
        'category', 'supplier'
    ).prefetch_related('location')

    if query:
        products = products.filter(
            Q(name__icontains=query) |
            Q(sku__icontains=query) |
            Q(location__icontains=query)
        )

    if category_id:
        products = products.filter(
            Q(location__id=category_id) | Q(category_id=category_id)
        ).distinct()
        area = Category.objects.filter(pk=category_id).first()
        products = order_queryset_by_area(products, area, product_id_field='pk', name_field='name')
    else:
        products = products.order_by('sequence_number', 'name')

    if supplier_id:
        products = products.filter(supplier_id=supplier_id)

    # Filtrado por estado de stock
    product_list_evaluated = []
    for p in products:
        if status_filter == 'out' and not p.is_out_of_stock:
            continue
        elif status_filter == 'low' and not p.is_below_minimum:
            continue
        elif status_filter == 'good' and p.stock_status['code'] not in ('good', 'regular'):
            continue
        product_list_evaluated.append(p)

    categories = Category.objects.all()
    suppliers = Supplier.objects.all()

    context = {
        'products': product_list_evaluated,
        'categories': categories,
        'suppliers': suppliers,
        'query': query,
        'selected_category': category_id,
        'selected_supplier': supplier_id,
        'selected_status': status_filter,
        'total_count': len(product_list_evaluated),
    }
    return render(request, 'stock/product_list.html', context)


def _area_order_rows(product=None):
    sequences = {}
    selected = set()
    fallback = 10
    if product is not None and product.pk:
        sequences = dict(product.area_sequences.values_list('area_id', 'sequence_number'))
        selected = set(product.location.values_list('id', flat=True))
        fallback = product.sequence_number or 10
    return [
        {
            'id': area.id,
            'name': area.name,
            'sequence': sequences.get(area.id, fallback),
            'selected': area.id in selected,
        }
        for area in Category.objects.order_by('name')
    ]


def _save_area_orders(product, post):
    selected_ids = set(product.location.values_list('id', flat=True))
    ProductAreaSequence.objects.filter(product=product).exclude(area_id__in=selected_ids).delete()
    for area_id in selected_ids:
        raw = str(post.get(f'area_order_{area_id}', '')).strip()
        if not raw:
            continue
        try:
            number = max(1, int(raw))
        except ValueError:
            continue
        ProductAreaSequence.objects.update_or_create(
            product=product,
            area_id=area_id,
            defaults={'sequence_number': number},
        )


@login_required
def product_create(request):
    """Creación de un nuevo producto."""
    if request.method == 'POST':
        form = ProductForm(request.POST)
        if form.is_valid():
            product = form.save()
            _save_area_orders(product, request.POST)
            initial_stock = product.current_stock or Decimal('0.00')
            if initial_stock > 0:
                StockMovement.objects.create(
                    product=product,
                    movement_type='entry',
                    quantity=initial_stock,
                    previous_stock=Decimal('0.00'),
                    new_stock=initial_stock,
                    reference="Inventario inicial al crear producto"
                )
            messages.success(request, f'Producto "{product.name}" creado con éxito.')
            return redirect('stock:product_list')
    else:
        form = ProductForm()

    return render(request, 'stock/product_form.html', {
        'form': form,
        'title': 'Nuevo Producto de Cocina / Barra',
        'area_orders': _area_order_rows(),
    })


@login_required
def product_edit(request, pk):
    """Edición de producto existente."""
    product = get_object_or_404(Product, pk=pk)
    if request.method == 'POST':
        form = ProductForm(request.POST, instance=product)
        if form.is_valid():
            product = form.save()
            _save_area_orders(product, request.POST)
            messages.success(request, f'Producto "{product.name}" actualizado correctamente.')
            return redirect('stock:product_list')
    else:
        form = ProductForm(instance=product)

    return render(request, 'stock/product_form.html', {
        'form': form,
        'title': f'Editar Producto: {product.name}',
        'product': product,
        'area_orders': _area_order_rows(product),
    })


@login_required
def product_detail(request, pk):
    """Ficha del producto con historial de movimientos (Kardex)."""
    product = get_object_or_404(
        Product.objects.select_related('category', 'supplier').prefetch_related('location'),
        pk=pk,
    )
    movements = product.movements.all()[:30]
    
    adjustment_form = QuickAdjustmentForm()

    context = {
        'product': product,
        'movements': movements,
        'adjustment_form': adjustment_form,
    }
    return render(request, 'stock/product_detail.html', context)


@login_required
def product_delete(request, pk):
    """Eliminación suave (desactivación) de un producto."""
    product = get_object_or_404(Product, pk=pk)
    if request.method == 'POST':
        product.is_active = False
        product.save()
        messages.success(request, f'Producto "{product.name}" archivado/eliminado correctamente.')
        return redirect('stock:product_list')
    return render(request, 'stock/product_confirm_delete.html', {'product': product})


@login_required
def product_adjust_stock(request, pk):
    """Ajuste rápido de inventario / Registro de merma o entrada de cocina."""
    product = get_object_or_404(Product, pk=pk)
    if request.method == 'POST':
        form = QuickAdjustmentForm(request.POST)
        if form.is_valid():
            mov_type = form.cleaned_data['movement_type']
            qty = form.cleaned_data['quantity']
            ref = form.cleaned_data['reference']

            prev_stock = product.current_stock or Decimal('0.00')
            if mov_type == 'waste':
                new_stock = max(Decimal('0.00'), prev_stock - qty)
                qty_diff = -qty
            elif mov_type == 'entry':
                new_stock = prev_stock + qty
                qty_diff = qty
            else: # adjustment directo
                new_stock = qty
                qty_diff = new_stock - prev_stock

            product.current_stock = new_stock
            product.save(update_fields=['current_stock', 'updated_at'])

            StockMovement.objects.create(
                product=product,
                movement_type=mov_type,
                quantity=qty_diff,
                previous_stock=prev_stock,
                new_stock=new_stock,
                reference=ref
            )
            messages.success(request, f'Stock de "{product.name}" actualizado a {new_stock} {product.get_unit_display()}.')
    return redirect('stock:product_detail', pk=pk)


# ==========================================
# REGISTRO DE INVENTARIO A LA FECHA (TOMA DE INVENTARIO)
# ==========================================

@login_required
def stock_audit_list(request):
    """Historial de tomas de inventario a la fecha."""
    audits = StockAudit.objects.select_related('category').all()
    return render(request, 'stock/audit_list.html', {'audits': audits})


@login_required
def stock_audit_create(request):
    """Inicia una nueva toma de inventario a una fecha determinada."""
    if request.method == 'POST':
        form = StockAuditForm(request.POST)
        if form.is_valid():
            audit = form.save()

            # Poblar automáticamente los productos a auditar ordenados por sequence_number
            created = _populate_audit_items(audit, audit.category)

            messages.success(request, f'Toma de inventario "{audit.title}" creada con {created} productos listos para conteo.')
            return redirect('stock:stock_audit_count', pk=audit.pk)
    else:
        today_str = timezone.now().strftime('%d/%m/%Y')
        initial_data = {
            'title': f'Toma de Inventario - {today_str}',
            'audit_date': timezone.now().date(),
            'location_area': 'Bodega Central',
        }
        form = StockAuditForm(initial=initial_data)

    return render(request, 'stock/audit_form.html', {'form': form, 'title': 'Iniciar Nueva Toma de Inventario'})


@login_required
def stock_audit_count(request, pk):
    """
    PANTALLA PRINCIPAL DE CONTEO Y CONTABILIZACIÓN:
    Permite buscar productos, ingresar la cantidad contada hasta la fecha,
    muestra cantidades mínimas, calcula automáticamente 'cantidades a pedir'
    y permite editarlas rápidamente.
    """
    audit = get_object_or_404(StockAudit.objects.select_related('category'), pk=pk)

    # Solo se cuentan los productos que tienen asignada el área de este conteo.
    area = audit.category
    if area is not None:
        _sync_area_items(audit, area)

    items = audit.items.select_related(
        'product', 'product__category', 'product__supplier'
    ).prefetch_related('product__location')
    if area is not None:
        items = items.filter(
            Q(product__location=area) | Q(product__category=area)
        ).distinct()
    items = order_queryset_by_area(items, area)

    if request.method == 'POST':
        action = request.POST.get('action', 'save_draft')

        # Procesar los valores enviados desde la tabla
        for item in items:
            counted_key = f'counted_{item.id}'
            checked_key = f'checked_{item.id}'
            order_key = f'order_{item.id}'
            notes_key = f'notes_{item.id}'

            _apply_item_capture(
                item,
                request.POST.get(counted_key, ''),
                request.POST.get(order_key, ''),
                request.POST.get(notes_key, ''),
                request.POST.get(checked_key) in ('on', 'true', '1', 'checked'),
            )

        if action == 'complete':
            if audit.session_id:
                # Área dentro de una sesión: se cierra el área y el stock se
                # aplicará de forma consolidada al finalizar la sesión completa.
                audit.mark_area_complete()
                messages.success(request, f'Área "{audit.location_area or audit.title}" finalizada. Al finalizar el conteo completo se consolidará el stock.')
                return redirect('stock:stock_count_session_detail', pk=audit.session_id)
            audit.apply_counts_to_inventory()
            messages.success(request, '¡Toma de inventario completada! Los stocks de los productos han sido actualizados en el sistema.')
            return redirect('stock:stock_audit_order_report', pk=audit.pk)
        else:
            messages.success(request, 'Progreso del inventario guardado correctamente en borrador.')
            return redirect('stock:stock_audit_count', pk=audit.pk)

    items = list(items)
    other_counts = _other_area_counts(audit, [item.product_id for item in items])
    for item in items:
        item.other_area_counts = other_counts.get(item.product_id, [])

    # Categorías presentes en la toma
    categories_present = set()
    for item in items:
        product_categories = {category.name for category in item.product.location.all()}
        if item.product.category:
            product_categories.add(item.product.category.name)
        categories_present.update(product_categories or {'Sin ubicación'})
    
    context = {
        'audit': audit,
        'items': items,
        'categories_present': sorted(list(categories_present)),
        'app_settings': AppSettings.load(),
    }
    return render(request, 'stock/audit_count.html', context)


@login_required
@require_POST
def stock_audit_item_save(request, pk, item_pk):
    """Guarda el progreso de un solo producto para poder continuar después."""
    audit = get_object_or_404(StockAudit, pk=pk)
    if audit.status == 'completed':
        return JsonResponse({'ok': False, 'error': 'Este conteo ya fue finalizado.'}, status=400)
    item = get_object_or_404(
        StockAuditItem.objects.select_related('product'),
        pk=item_pk,
        audit=audit,
    )
    _apply_item_capture(
        item,
        request.POST.get('counted', ''),
        request.POST.get('order', ''),
        request.POST.get('notes', ''),
        request.POST.get('checked') in ('on', 'true', '1', 'checked'),
    )
    return JsonResponse({'ok': True, 'item': _item_capture_payload(item)})


@login_required
@require_POST
def stock_audit_mark_pending_zero(request, pk):
    """Coloca en cero los productos pendientes indicados y guarda ese progreso."""
    audit = get_object_or_404(StockAudit, pk=pk)
    if audit.status == 'completed':
        return JsonResponse({'ok': False, 'error': 'Este conteo ya fue finalizado.'}, status=400)

    raw_ids = request.POST.getlist('item_ids')
    ids = []
    for raw_id in raw_ids:
        try:
            ids.append(int(raw_id))
        except (TypeError, ValueError):
            continue

    results = []
    items = audit.items.select_related('product').filter(id__in=ids)
    for item in items:
        _apply_item_capture(item, '0', '', item.notes, is_checked=False)
        results.append(_item_capture_payload(item))
    return JsonResponse({'ok': True, 'updated': len(results), 'items': results})


@login_required
def stock_audit_order_report(request, pk):
    """
    REPORTE DE SOLICITUD DE PRODUCTOS / PEDIDOS DE COMPRA:
    Muestra los productos que requieren reposición a partir de la toma de inventario,
    agrupados por proveedor y ordenados por su ordinal/secuencia.
    """
    audit = get_object_or_404(StockAudit, pk=pk)
    
    order_items = audit.items.filter(final_order__gt=0).select_related(
        'product', 'product__category', 'product__supplier'
    ).prefetch_related('product__location')
    order_items = order_queryset_by_area(order_items, audit.category)

    # Agrupar por Proveedor
    suppliers_dict = {}
    for item in order_items:
        supp = item.product.supplier
        supp_key = supp.id if supp else 'sin_proveedor'
        if supp_key not in suppliers_dict:
            suppliers_dict[supp_key] = {
                'supplier': supp,
                'items': [],
                'total_cost': Decimal('0.00')
            }
        suppliers_dict[supp_key]['items'].append(item)
        suppliers_dict[supp_key]['total_cost'] += item.subtotal_cost

    total_order_cost = sum(group['total_cost'] for group in suppliers_dict.values())
    total_order_items = order_items.count()

    context = {
        'audit': audit,
        'order_items': order_items,
        'suppliers_groups': suppliers_dict.values(),
        'total_order_cost': total_order_cost,
        'total_order_items': total_order_items,
    }
    return render(request, 'stock/audit_order_report.html', context)


@login_required
def stock_audit_export_csv(request, pk):
    """Descarga de reporte de solicitud en formato CSV ordenado por ordinal/secuencia."""
    audit = get_object_or_404(StockAudit, pk=pk)
    order_items = order_queryset_by_area(
        audit.items.filter(final_order__gt=0).select_related('product', 'product__supplier'),
        audit.category,
    )

    response = HttpResponse(content_type='text/csv; charset=utf-8')
    filename = f"solicitud_compra_{audit.audit_date.strftime('%Y%m%d')}_{audit.id}.csv"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    writer = csv.writer(response)
    writer.writerow(['Secuencia', 'Producto', 'Proveedor', 'Cantidad a Pedir', 'Unidad', 'Stock Contado', 'Stock Mínimo', 'Costo Est. Unitario (Q)', 'Subtotal Estimado (Q)', 'Notas'])

    for item in order_items:
        writer.writerow([
            getattr(item, 'area_sequence', item.product.sequence_number),
            item.product.name,
            item.product.supplier.name if item.product.supplier else 'Sin Proveedor',
            f"{item.final_order:.2f}",
            item.product.get_unit_display(),
            f"{item.counted_stock:.2f}" if item.counted_stock is not None else "0.00",
            f"{item.min_stock:.2f}",
            f"{(item.product.cost_price or Decimal('0.00')):.2f}",
            f"{item.subtotal_cost:.2f}",
            item.notes
        ])

    return response


@login_required
def stock_audit_delete(request, pk):
    """Eliminar una toma de inventario."""
    audit = get_object_or_404(StockAudit, pk=pk)
    if request.method == 'POST':
        audit.delete()
        messages.success(request, 'Registro de inventario eliminado.')
        return redirect('stock:stock_audit_list')
    return render(request, 'stock/audit_confirm_delete.html', {'audit': audit})


# ==========================================
# REGISTRO DE CONTEO DE STOCK (SESIÓN MULTI-ÁREA)
# ==========================================

@login_required
def stock_count_session_list(request):
    """Historial de registros de conteo de stock (sesiones con múltiples áreas)."""
    sessions = StockCountSession.objects.all()
    return render(request, 'stock/session_list.html', {'sessions': sessions})


def _create_area_audit(session, area):
    """Crea el conteo (StockAudit) de un área/ubicación dentro de una sesión."""
    audit = StockAudit.objects.create(
        session=session,
        category=area,
        location_area=area.name,
        title=f"{session.title} - {area.name}",
        audit_date=session.count_date,
        responsible=session.responsible,
        notes=session.notes,
        status='draft',
    )
    _populate_audit_items(audit, area)
    return audit


@login_required
def stock_count_session_create(request):
    """Inicia un nuevo registro de conteo de stock seleccionando una o varias áreas."""
    if request.method == 'POST':
        form = StockCountSessionForm(request.POST)
        if form.is_valid():
            session = form.save()
            areas = form.cleaned_data['areas']
            for area in areas:
                _create_area_audit(session, area)
            messages.success(request, f'Registro de conteo "{session.title}" creado con {areas.count()} área(s) lista(s) para contar.')
            return redirect('stock:stock_count_session_detail', pk=session.pk)
    else:
        today_str = timezone.now().strftime('%d/%m/%Y')
        form = StockCountSessionForm(initial={
            'title': f'Conteo de Stock - {today_str}',
            'count_date': timezone.now().date(),
        })
    return render(request, 'stock/session_form.html', {'form': form, 'title': 'Nuevo Registro de Conteo de Stock'})


@login_required
def stock_count_session_detail(request, pk):
    """Panel del registro de conteo: lista las áreas, su avance y permite ir contando."""
    session = get_object_or_404(StockCountSession, pk=pk)
    areas = session.audits.select_related('category').annotate(
        n_items=Count('items', distinct=True),
        n_counted=Count('items', filter=Q(items__counted_stock__isnull=False), distinct=True),
        n_to_order=Count('items', filter=Q(items__final_order__gt=0), distinct=True),
    ).order_by('category__name', 'id')

    existing_ids = list(session.audits.values_list('category_id', flat=True))
    available_areas = Category.objects.exclude(id__in=existing_ids)

    context = {
        'session': session,
        'areas': areas,
        'available_areas': available_areas,
    }
    return render(request, 'stock/session_detail.html', context)


@login_required
def stock_count_session_add_area(request, pk):
    """Agrega un área/ubicación adicional a una sesión en proceso."""
    session = get_object_or_404(StockCountSession, pk=pk)
    if request.method != 'POST':
        return redirect('stock:stock_count_session_detail', pk=pk)
    if session.status == 'completed':
        messages.error(request, 'No puedes agregar áreas a un conteo ya finalizado.')
        return redirect('stock:stock_count_session_detail', pk=pk)

    area = get_object_or_404(Category, pk=request.POST.get('area'))
    if session.audits.filter(category=area).exists():
        messages.warning(request, f'El área "{area.name}" ya está incluida en este conteo.')
    else:
        _create_area_audit(session, area)
        messages.success(request, f'Área "{area.name}" agregada al conteo.')
    return redirect('stock:stock_count_session_detail', pk=pk)


@login_required
def stock_count_session_remove_area(request, pk, audit_pk):
    """Quita un área de la sesión (solo si la sesión sigue en proceso)."""
    session = get_object_or_404(StockCountSession, pk=pk)
    audit = get_object_or_404(StockAudit, pk=audit_pk, session=session)
    if request.method == 'POST':
        if session.status == 'completed':
            messages.error(request, 'No puedes modificar un conteo ya finalizado.')
        else:
            area_name = audit.location_area or audit.title
            audit.delete()
            messages.success(request, f'Área "{area_name}" quitada del conteo.')
    return redirect('stock:stock_count_session_detail', pk=pk)


@login_required
def stock_count_session_reopen_area(request, pk, audit_pk):
    """Reabre un área finalizada para seguir contando (solo si la sesión sigue en proceso)."""
    session = get_object_or_404(StockCountSession, pk=pk)
    audit = get_object_or_404(StockAudit, pk=audit_pk, session=session)
    if request.method == 'POST':
        if session.status == 'completed':
            messages.error(request, 'El conteo ya fue finalizado y aplicado al inventario.')
        else:
            audit.reopen_area()
            messages.success(request, f'Área "{audit.location_area or audit.title}" reabierta para continuar contando.')
    return redirect('stock:stock_count_session_detail', pk=pk)


def _area_column_title(audit):
    return audit.location_area or (audit.category.name if audit.category_id else audit.title) or 'Área'


def _count_display_format(value, decimal_separator):
    """Muestra 5.00 o 5,00 sin mover el número que usa la suma.

    El separador va entre comillas para que Excel lo pinte tal cual. Si la coma
    queda suelta en el formato, Excel la toma como miles y un 5 se ve como 0,5.
    """
    text = f'{Decimal(value).quantize(Decimal("0.01")):.2f}'
    if decimal_separator == ',':
        text = text.replace('.', ',')
    return f'"{text}"'


def build_session_count_workbook(session, audits, decimal_separator='.'):
    """Arma el Excel del conteo: una fila por producto y una columna por área."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    items = StockAuditItem.objects.filter(audit__in=audits).select_related('product')
    products = {}
    for item in items:
        entry = products.setdefault(
            item.product_id,
            {'product': item.product, 'counts': {}, 'orders': {}},
        )
        entry['counts'][item.audit_id] = item.counted_stock
        entry['orders'][item.audit_id] = item.final_order or Decimal('0.00')
    first_area = audits[0].category if audits else None
    area_sequences = sequences_for_area(first_area, products.keys())

    def _report_sequence(product):
        return area_sequences.get(product.id, product.sequence_number)

    ordered = sorted(
        products.values(),
        key=lambda entry: (_report_sequence(entry['product']), entry['product'].name.lower()),
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Conteo'

    fixed_headers = ['Orden', 'Grupo', 'Código', 'Nombre del producto', 'Presentación']
    area_headers = [_area_column_title(audit) for audit in audits]
    headers = fixed_headers + area_headers + ['Total existencias', 'Pedido']
    last_col = len(headers)
    first_area_col = len(fixed_headers) + 1
    last_area_col = first_area_col + len(area_headers) - 1
    total_col = last_col - 1
    order_col = last_col

    title_font = Font(name='Calibri', bold=True, size=18, color='1F2933')
    header_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='0F7B6C')
    title_fill = PatternFill('solid', fgColor='F4FBF8')
    thin = Border(
        left=Side(style='thin', color='D9E2DF'),
        right=Side(style='thin', color='D9E2DF'),
        top=Side(style='thin', color='D9E2DF'),
        bottom=Side(style='thin', color='D9E2DF'),
    )
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center')

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(fixed_headers))
    sheet['A1'] = session.title
    sheet['A1'].font = title_font
    sheet['A1'].alignment = Alignment(horizontal='left', vertical='center')
    sheet['A1'].fill = title_fill

    if area_headers:
        sheet.merge_cells(start_row=1, start_column=first_area_col, end_row=1, end_column=last_col)
        period_cell = sheet.cell(1, first_area_col, session.count_date.strftime('%d/%m/%Y'))
        period_cell.font = title_font
        period_cell.alignment = Alignment(horizontal='center', vertical='center')
        period_cell.fill = title_fill
    for col in range(1, last_col + 1):
        sheet.cell(1, col).fill = title_fill
        sheet.cell(1, col).border = thin

    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(2, col, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin

    for row_index, entry in enumerate(ordered, start=3):
        product = entry['product']
        values = [
            _report_sequence(product),
            product.group or '',
            product.sku or '',
            product.name,
            product.presentation or '',
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, col, value)
            cell.border = thin
            cell.alignment = center if col == 1 else left_align
            if col == 1:
                cell.font = Font(name='Calibri', bold=True)

        row_total = Decimal('0.00')
        for offset, audit in enumerate(audits):
            counted = entry['counts'].get(audit.id)
            amount = Decimal('0.00') if counted is None else counted
            row_total += amount
            col = first_area_col + offset
            cell = sheet.cell(row_index, col, amount)
            cell.border = thin
            cell.alignment = center
            cell.number_format = _count_display_format(amount, decimal_separator)

        start_letter = get_column_letter(first_area_col)
        end_letter = get_column_letter(last_area_col)
        total_cell = sheet.cell(
            row_index,
            total_col,
            f'=SUM({start_letter}{row_index}:{end_letter}{row_index})',
        )
        total_cell.border = thin
        total_cell.alignment = center
        total_cell.font = Font(name='Calibri', bold=True)
        total_cell.number_format = _count_display_format(row_total, decimal_separator)

        order_total = sum(
            (entry['orders'].get(audit.id, Decimal('0.00')) for audit in audits),
            Decimal('0.00'),
        )
        order_cell = sheet.cell(row_index, order_col, order_total if order_total > 0 else None)
        order_cell.border = thin
        order_cell.alignment = center
        if order_total > 0:
            order_cell.number_format = _count_display_format(order_total, decimal_separator)

    widths = [12, 18, 16, 36, 22] + [16] * len(area_headers) + [20, 16]
    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 22
    sheet.freeze_panes = 'A3'
    sheet.auto_filter.ref = f'A2:{get_column_letter(last_col)}{max(2, 2 + len(ordered))}'
    return workbook


@login_required
def stock_count_session_export_excel(request, pk):
    """Descarga el conteo de las áreas elegidas, una columna por área."""
    session = get_object_or_404(StockCountSession, pk=pk)
    if request.method != 'POST':
        return redirect('stock:stock_count_session_detail', pk=pk)

    raw_ids = request.POST.getlist('areas')
    ids = []
    for raw_id in raw_ids:
        try:
            ids.append(int(raw_id))
        except (TypeError, ValueError):
            continue
    audits = list(
        session.audits.filter(id__in=ids).select_related('category').order_by('category__name', 'id')
    )
    if not audits:
        messages.warning(request, 'Selecciona al menos un área para generar el Excel.')
        return redirect('stock:stock_count_session_detail', pk=pk)

    decimal_separator = ',' if request.POST.get('decimal_separator') == ',' else '.'
    workbook = build_session_count_workbook(session, audits, decimal_separator)
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    slug = ''.join(ch if ch.isalnum() else '_' for ch in session.title).strip('_') or f'conteo_{session.pk}'
    response['Content-Disposition'] = f'attachment; filename="{slug}.xlsx"'
    workbook.save(response)
    return response


@login_required
def stock_count_session_finalize(request, pk):
    """Finaliza la sesión: consolida las áreas y aplica el stock al inventario."""
    session = get_object_or_404(StockCountSession, pk=pk)
    if request.method == 'POST':
        if session.apply_consolidated_to_inventory():
            messages.success(request, f'Conteo "{session.title}" finalizado y consolidado. Los stocks fueron actualizados en el sistema.')
        else:
            messages.warning(request, 'Este conteo ya había sido finalizado.')
    return redirect('stock:stock_count_session_detail', pk=pk)


# ==========================================
# SOLICITUDES Y ÓRDENES DE COMPRA (PEDIDOS)
# ==========================================

@login_required
def purchase_order_list(request):
    """Listado de pedidos y solicitudes de compra emitidas."""
    orders = PurchaseOrder.objects.select_related('supplier', 'audit').all()
    return render(request, 'stock/purchase_order_list.html', {'orders': orders})


@login_required
def create_orders_from_audit(request, pk):
    """
    Genera formalmente una o varias órdenes de compra a partir de la toma de inventario:
    Crea órdenes separadas por proveedor o una orden global según prefiera el usuario.
    """
    audit = get_object_or_404(StockAudit, pk=pk)
    order_items = order_queryset_by_area(
        audit.items.filter(final_order__gt=0).select_related('product', 'product__supplier'),
        audit.category,
    )

    if not order_items.exists():
        messages.warning(request, 'No hay productos con cantidades a pedir en esta toma de inventario.')
        return redirect('stock:stock_audit_order_report', pk=pk)

    mode = request.POST.get('mode', 'by_supplier') # 'by_supplier' o 'single'
    created_count = 0

    if mode == 'by_supplier':
        # Agrupar por proveedor
        grouped = {}
        for item in order_items:
            supp = item.product.supplier
            grouped.setdefault(supp, []).append(item)

        for supp, items_list in grouped.items():
            now_code = timezone.now().strftime('%y%m%d%H%M')
            supp_suffix = supp.name[:3].upper() if supp else 'GEN'
            code = f"PED-{now_code}-{supp_suffix}"

            po = PurchaseOrder.objects.create(
                code=code,
                audit=audit,
                supplier=supp,
                date=audit.audit_date,
                responsible=audit.responsible,
                notes=f"Generado automáticamente desde la Toma de Inventario: {audit.title}"
            )
            for it in items_list:
                PurchaseOrderItem.objects.create(
                    order=po,
                    product=it.product,
                    quantity=it.final_order,
                    unit=it.product.unit,
                    unit_price=it.product.cost_price or Decimal('0.00'),
                    notes=it.notes
                )
            created_count += 1

    else:
        # Una sola orden global
        now_code = timezone.now().strftime('%y%m%d%H%M')
        code = f"PED-{now_code}-ALL"
        po = PurchaseOrder.objects.create(
            code=code,
            audit=audit,
            date=audit.audit_date,
            responsible=audit.responsible,
            notes=f"Solicitud consolidada desde: {audit.title}"
        )
        for it in order_items:
            PurchaseOrderItem.objects.create(
                order=po,
                product=it.product,
                quantity=it.final_order,
                unit=it.product.unit,
                unit_price=it.product.cost_price or Decimal('0.00'),
                notes=it.notes
            )
        created_count = 1

    messages.success(request, f'Se han generado {created_count} Solicitud(es) de Compra con éxito.')
    return redirect('stock:purchase_order_list')


@login_required
def purchase_order_detail(request, pk):
    """Vista detallada de una orden de compra."""
    order = get_object_or_404(PurchaseOrder.objects.select_related('supplier', 'audit'), pk=pk)
    items = order.items.select_related('product', 'product__category').prefetch_related(
        'product__location'
    ).order_by('product__sequence_number', 'product__name')

    # Pre-generar texto para compartir por WhatsApp
    restaurant_name = "Mi Restaurante"
    wa_lines = [
        f"📋 *SOLICITUD DE PEDIDO - {restaurant_name}*",
        f"📅 Fecha: {order.date.strftime('%d/%m/%Y')}",
        f"🔖 Ref: {order.code}",
    ]
    if order.supplier:
        wa_lines.append(f"🏢 Proveedor: {order.supplier.name}")
    wa_lines.append("\n*PRODUCTOS SOLICITADOS:*")
    
    for it in items:
        wa_lines.append(f"• {it.quantity} {it.product.get_unit_display()} - {it.product.name}")
        if it.notes:
            wa_lines.append(f"   _(Nota: {it.notes})_")

    wa_text = "\n".join(wa_lines)

    context = {
        'order': order,
        'items': items,
        'wa_text': wa_text,
    }
    return render(request, 'stock/purchase_order_detail.html', context)


@login_required
def purchase_order_toggle_lock(request, pk):
    """Bloquear o desbloquear un pedido ya realizado."""
    order = get_object_or_404(PurchaseOrder, pk=pk)
    if request.method == 'POST':
        if order.is_locked:
            order.is_locked = False
            order.locked_at = None
            order.locked_by = ''
            order.save()
            messages.success(request, f'Pedido {order.code} DESBLOQUEADO. Ahora permite modificaciones.')
        else:
            order.is_locked = True
            order.locked_at = timezone.now()
            order.locked_by = request.user.username if request.user.is_authenticated else 'Usuario'
            order.save()
            messages.success(request, f'Pedido {order.code} BLOQUEADO exitosamente contra modificaciones.')
    return redirect('stock:purchase_order_detail', pk=pk)


@login_required
def purchase_order_change_status(request, pk):
    """Cambiar estado de pedido (Pendiente -> Enviado -> Recibido). Si está bloqueado no lo permite."""
    order = get_object_or_404(PurchaseOrder, pk=pk)
    if order.is_locked:
        messages.error(request, f'El pedido {order.code} está BLOQUEADO. Desbloquéalo primero para cambiar su estado.')
        return redirect('stock:purchase_order_detail', pk=pk)

    if request.method == 'POST':
        new_status = request.POST.get('status')
        if new_status in dict(PurchaseOrder.STATUS_CHOICES):
            prev_status = order.status
            order.status = new_status
            order.save()

            # Si pasa a 'received' y antes no lo estaba, ingresar al stock
            if new_status == 'received' and prev_status != 'received':
                for it in order.items.all():
                    prod = it.product
                    prev_stock = prod.current_stock or Decimal('0.00')
                    prod.current_stock = prev_stock + it.quantity
                    prod.save(update_fields=['current_stock', 'updated_at'])

                    StockMovement.objects.create(
                        product=prod,
                        movement_type='entry',
                        quantity=it.quantity,
                        previous_stock=prev_stock,
                        new_stock=prod.current_stock,
                        reference=f"Recepción de Pedido {order.code}"
                    )
                messages.success(request, f'Pedido {order.code} marcado como RECIBIDO. Stock ingresado al almacén.')
            else:
                messages.success(request, f'Estado del pedido {order.code} actualizado a "{order.get_status_display()}".')

    return redirect('stock:purchase_order_detail', pk=pk)


# ==========================================
# REPORTE GENERAL DIARIO & CONSOLIDACIÓN MULTI-REGISTRO
# ==========================================

@login_required
def consolidated_selection(request):
    """
    Pantalla de selección de múltiples tomas de inventario (ej. Bodega + Barra + Cocina)
    para generar un reporte general diario consolidado de pedidos.
    """
    date_filter = request.GET.get('date', '')
    query = request.GET.get('q', '').strip()

    audits = StockAudit.objects.select_related('category').all()

    if date_filter:
        audits = audits.filter(audit_date=date_filter)

    if query:
        audits = audits.filter(Q(title__icontains=query) | Q(location_area__icontains=query) | Q(responsible__icontains=query))

    today_str = timezone.now().date().isoformat()

    context = {
        'audits': audits,
        'today_str': today_str,
        'date_filter': date_filter,
        'query': query,
    }
    return render(request, 'stock/consolidated_selection.html', context)


@login_required
def consolidated_report_view(request):
    """
    REPORTE GENERAL DIARIO CONSOLIDADO:
    Consolida las cantidades contadas del mismo insumo en distintas ubicaciones
    (ej: Bodega, Barra, Cocina) y calcula el faltante total neto para pedir.
    """
    if request.method == 'POST':
        audit_ids = request.POST.getlist('audit_ids')
    else:
        audit_ids_param = request.GET.get('audits', '')
        audit_ids = [aid for aid in audit_ids_param.split(',') if aid.strip().isdigit()]

    if not audit_ids:
        messages.warning(request, 'Debes seleccionar al menos un registro de conteo para consolidar.')
        return redirect('stock:consolidated_selection')

    selected_audits = StockAudit.objects.filter(id__in=audit_ids).order_by('audit_date', 'title')

    # Obtener todos los items contados de esas tomas
    audit_items = StockAuditItem.objects.filter(
        audit__in=selected_audits,
        counted_stock__isnull=False
    ).select_related('product', 'product__category', 'product__supplier', 'audit').order_by(
        'product__sequence_number', 'product__name'
    )

    # Agrupar por producto
    products_map = {}
    for item in audit_items:
        prod = item.product
        pid = prod.id
        if pid not in products_map:
            products_map[pid] = {
                'product': prod,
                'locations': [],
                'total_counted': Decimal('0.00'),
                'min_stock': prod.min_stock or Decimal('0.00'),
                'ideal_stock': prod.ideal_stock or Decimal('0.00'),
                'suggested_order': Decimal('0.00'),
                'final_order': Decimal('0.00'),
                'notes_list': [],
            }

        loc_label = item.audit.location_area or item.audit.title
        products_map[pid]['locations'].append({
            'audit_title': item.audit.title,
            'location': loc_label,
            'counted': item.counted_stock,
            'unit': prod.get_unit_display()
        })
        products_map[pid]['total_counted'] += item.counted_stock
        if item.notes:
            products_map[pid]['notes_list'].append(f"{loc_label}: {item.notes}")

    # Calcular sugerencia neta de cada producto
    consolidated_items = []
    for pid, data in products_map.items():
        prod = data['product']
        ideal_stock = prod.ideal_stock or Decimal('0.00')
        min_stock = prod.min_stock or Decimal('0.00')
        target = ideal_stock if ideal_stock > 0 else min_stock
        suggested = max(Decimal('0.00'), target - data['total_counted'])
        data['suggested_order'] = suggested
        data['final_order'] = suggested
        data['subtotal_cost'] = suggested * (prod.cost_price or Decimal('0.00'))
        consolidated_items.append(data)

    first_area = None
    audits_by_id = {str(audit.pk): audit for audit in selected_audits}
    for raw_id in audit_ids:
        audit = audits_by_id.get(str(raw_id))
        if audit is not None:
            first_area = audit.category
            break
    area_sequences = sequences_for_area(first_area, products_map.keys())

    def _area_seq(product):
        return area_sequences.get(product.id, product.sequence_number)

    for data in consolidated_items:
        data['area_sequence'] = _area_seq(data['product'])
    consolidated_items.sort(
        key=lambda x: (x['area_sequence'], x['product'].name.lower())
    )

    # Filtrar aquellos que requieren pedido
    items_to_order = [it for it in consolidated_items if it['final_order'] > 0]

    # Agrupar por proveedor
    suppliers_dict = {}
    for it in items_to_order:
        supp = it['product'].supplier
        skey = supp.id if supp else 'sin_proveedor'
        if skey not in suppliers_dict:
            suppliers_dict[skey] = {
                'supplier': supp,
                'items': [],
                'total_cost': Decimal('0.00')
            }
        suppliers_dict[skey]['items'].append(it)
        suppliers_dict[skey]['total_cost'] += it['subtotal_cost']

    total_order_cost = sum(group['total_cost'] for group in suppliers_dict.values())
    total_items_to_order = len(items_to_order)

    audit_ids_str = ','.join(map(str, audit_ids))

    context = {
        'selected_audits': selected_audits,
        'consolidated_items': consolidated_items,
        'items_to_order': items_to_order,
        'suppliers_groups': suppliers_dict.values(),
        'total_order_cost': total_order_cost,
        'total_items_to_order': total_items_to_order,
        'audit_ids_str': audit_ids_str,
    }
    return render(request, 'stock/consolidated_report.html', context)


@login_required
def consolidated_export_csv(request):
    """Exportar reporte consolidado a CSV."""
    audit_ids_param = request.GET.get('audits', '')
    audit_ids = [aid for aid in audit_ids_param.split(',') if aid.strip().isdigit()]

    if not audit_ids:
        return HttpResponse('No audits specified', status=400)

    selected_audits = StockAudit.objects.filter(id__in=audit_ids)
    audit_items = StockAuditItem.objects.filter(
        audit__in=selected_audits,
        counted_stock__isnull=False
    ).select_related('product', 'product__supplier')

    products_map = {}
    for item in audit_items:
        prod = item.product
        if prod.id not in products_map:
            products_map[prod.id] = {
                'product': prod,
                'total_counted': Decimal('0.00'),
                'locations': []
            }
        loc_label = item.audit.location_area or item.audit.title
        products_map[prod.id]['total_counted'] += item.counted_stock
        products_map[prod.id]['locations'].append(f"{loc_label}: {item.counted_stock}")

    response = HttpResponse(content_type='text/csv; charset=utf-8')
    filename = f"reporte_consolidado_{timezone.now().strftime('%Y%m%d_%H%M')}.csv"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    writer = csv.writer(response)
    writer.writerow(['Secuencia', 'Producto', 'Proveedor', 'Total Contado Consolidado', 'Desglose por Ubicación', 'Stock Mínimo', 'Cantidad a Pedir', 'Unidad', 'Costo Unit. (Q)', 'Subtotal Est. (Q)'])

    first_audit = StockAudit.objects.filter(id__in=audit_ids).select_related('category').order_by('id').first()
    if audit_ids:
        preferred = StockAudit.objects.filter(id=audit_ids[0]).select_related('category').first()
        if preferred is not None:
            first_audit = preferred
    csv_sequences = sequences_for_area(
        first_audit.category if first_audit else None,
        products_map.keys(),
    )
    for pid, data in sorted(
        products_map.items(),
        key=lambda x: (
            csv_sequences.get(x[1]['product'].id, x[1]['product'].sequence_number),
            x[1]['product'].name.lower(),
        ),
    ):
        prod = data['product']
        ideal_stock = prod.ideal_stock or Decimal('0.00')
        min_stock = prod.min_stock or Decimal('0.00')
        target = ideal_stock if ideal_stock > 0 else min_stock
        order_qty = max(Decimal('0.00'), target - data['total_counted'])
        if order_qty > 0:
            writer.writerow([
                csv_sequences.get(prod.id, prod.sequence_number),
                prod.name,
                prod.supplier.name if prod.supplier else 'Sin Proveedor',
                f"{data['total_counted']:.2f}",
                " | ".join(data['locations']),
                f"{min_stock:.2f}",
                f"{order_qty:.2f}",
                prod.get_unit_display(),
                f"{(prod.cost_price or Decimal('0.00')):.2f}",
                f"{(order_qty * (prod.cost_price or Decimal('0.00'))):.2f}"
            ])

    return response


@login_required
def create_orders_from_consolidated(request):
    """Generar órdenes de compra a partir del reporte diario consolidado."""
    audit_ids_param = request.POST.get('audits', '')
    audit_ids = [aid for aid in audit_ids_param.split(',') if aid.strip().isdigit()]

    if not audit_ids:
        messages.error(request, 'No se especificaron tomas para crear pedidos.')
        return redirect('stock:consolidated_selection')

    selected_audits = StockAudit.objects.filter(id__in=audit_ids)
    audit_items = StockAuditItem.objects.filter(
        audit__in=selected_audits,
        counted_stock__isnull=False
    ).select_related('product', 'product__supplier')

    products_map = {}
    for item in audit_items:
        prod = item.product
        if prod.id not in products_map:
            products_map[prod.id] = {
                'product': prod,
                'total_counted': Decimal('0.00'),
            }
        products_map[prod.id]['total_counted'] += item.counted_stock

    # Agrupar por proveedor los insumos que requieren reposición
    supp_orders = {}
    for pid, data in products_map.items():
        prod = data['product']
        ideal_stock = prod.ideal_stock or Decimal('0.00')
        min_stock = prod.min_stock or Decimal('0.00')
        target = ideal_stock if ideal_stock > 0 else min_stock
        order_qty = max(Decimal('0.00'), target - data['total_counted'])
        if order_qty > 0:
            supp = prod.supplier
            supp_orders.setdefault(supp, []).append((prod, order_qty))

    if not supp_orders:
        messages.warning(request, 'No hay insumos que requieran pedido en esta consolidación.')
        return redirect('stock:consolidated_selection')

    created_count = 0
    now_str = timezone.now().strftime('%y%m%d%H%M')
    for supp, items_list in supp_orders.items():
        s_code = supp.name[:3].upper() if supp else 'GEN'
        code = f"PED-CONS-{now_str}-{s_code}"
        po = PurchaseOrder.objects.create(
            code=code,
            supplier=supp,
            date=timezone.now().date(),
            responsible=request.user.username if request.user.is_authenticated else 'Consolidado',
            notes=f"Generado a partir de {selected_audits.count()} tomas de inventario consolidadas."
        )
        for prod, qty in items_list:
            PurchaseOrderItem.objects.create(
                order=po,
                product=prod,
                quantity=qty,
                unit=prod.unit,
                unit_price=prod.cost_price or Decimal('0.00'),
            )
        created_count += 1

    messages.success(request, f'Se crearon exitosamente {created_count} pedidos de compra consolidados.')
    return redirect('stock:purchase_order_list')


# ==========================================
# REPORTE RÁPIDO DE COMPRAS EXPRESS
# ==========================================

@login_required
def quick_order_report(request):
    """
    Genera un reporte express de reposición inmediata
    basado en los productos activos que actualmente están bajo mínimo.
    """
    low_products = [p for p in Product.objects.filter(is_active=True).select_related('category', 'supplier').order_by('sequence_number', 'category__name', 'name') if p.is_below_minimum]
    
    context = {
        'low_products': low_products,
        'today': timezone.now().date(),
    }
    return render(request, 'stock/quick_order_report.html', context)


# ==========================================
# GESTIÓN DE PROVEEDORES Y CATEGORÍAS
# ==========================================

@login_required
def supplier_list(request):
    suppliers = Supplier.objects.all()
    return render(request, 'stock/supplier_list.html', {'suppliers': suppliers})


@login_required
def supplier_create(request):
    if request.method == 'POST':
        form = SupplierForm(request.POST)
        if form.is_valid():
            supp = form.save()
            messages.success(request, f'Proveedor "{supp.name}" registrado.')
            return redirect('stock:supplier_list')
    else:
        form = SupplierForm()
    return render(request, 'stock/supplier_form.html', {'form': form, 'title': 'Nuevo Proveedor'})


@login_required
def supplier_edit(request, pk):
    supp = get_object_or_404(Supplier, pk=pk)
    if request.method == 'POST':
        form = SupplierForm(request.POST, instance=supp)
        if form.is_valid():
            form.save()
            messages.success(request, f'Proveedor "{supp.name}" actualizado.')
            return redirect('stock:supplier_list')
    else:
        form = SupplierForm(instance=supp)
    return render(request, 'stock/supplier_form.html', {'form': form, 'title': f'Editar: {supp.name}', 'supplier': supp})


@login_required
def category_list(request):
    categories = Category.objects.all()
    return render(request, 'stock/category_list.html', {'categories': categories})


@login_required
def category_create(request):
    if request.method == 'POST':
        form = CategoryForm(request.POST)
        if form.is_valid():
            cat = form.save()
            messages.success(request, f'Área/Categoría "{cat.name}" creada.')
            return redirect('stock:category_list')
    else:
        form = CategoryForm()
    return render(request, 'stock/category_form.html', {'form': form, 'title': 'Nueva Área / Categoría'})


@login_required
def category_edit(request, pk):
    cat = get_object_or_404(Category, pk=pk)
    if request.method == 'POST':
        form = CategoryForm(request.POST, instance=cat)
        if form.is_valid():
            form.save()
            messages.success(request, f'Categoría "{cat.name}" actualizada.')
            return redirect('stock:category_list')
    else:
        form = CategoryForm(instance=cat)
    return render(request, 'stock/category_form.html', {'form': form, 'title': f'Editar: {cat.name}'})


# ==========================================
# MOVIMIENTOS / KARDEX
# ==========================================

@login_required
def movement_list(request):
    movements = StockMovement.objects.select_related('product', 'product__category').all()[:100]
    return render(request, 'stock/movement_list.html', {'movements': movements})


# ==========================================
# IMPORTACIÓN DE PRODUCTOS DESDE EXCEL
# ==========================================

@login_required
def product_import(request):
    if request.method == 'POST':
        upload = request.FILES.get('file')
        if not upload:
            messages.error(request, 'Selecciona un archivo Excel (.xlsx) para importar.')
        else:
            try:
                summary = import_products_from_workbook(upload)
            except Exception as exc:  # noqa: BLE001
                messages.error(request, f'No se pudo leer el archivo: {exc}')
            else:
                messages.success(
                    request,
                    'Importación completada: '
                    f"{summary['created']} creados, {summary['updated']} actualizados, "
                    f"{summary['skipped']} omitidos, {summary['areas_created']} áreas nuevas."
                )
                for error in summary['errors'][:10]:
                    messages.warning(request, error)
                return redirect('stock:product_list')
    return render(request, 'stock/product_import.html', {
        'title': 'Importar Productos desde Excel',
    })
