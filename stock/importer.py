import re

from django.db import transaction

from .models import Category, Product


HEADER_MAP = {
    'sequence_number': 'sequence_number',
    'secuencia': 'sequence_number',
    'numero': 'sequence_number',
    'no': 'sequence_number',
    'ordinal': 'sequence_number',
    'sku': 'sku',
    'codigo': 'sku',
    'código': 'sku',
    'code': 'sku',
    'reference': 'sku',
    'referencia': 'sku',
    'group': 'group',
    'grupo': 'group',
    'name': 'name',
    'nombre': 'name',
    'producto': 'name',
    'insumo': 'name',
    'nombre_producto': 'name',
    'nombre_del_producto': 'name',
    'presentation': 'presentation',
    'presentacion': 'presentation',
    'presentación': 'presentation',
    'location': 'location',
    'ubicacion': 'location',
    'ubicación': 'location',
    'ubicaciones': 'location',
    'area': 'location',
    'areas': 'location',
    'seccion': 'location',
    'sección': 'location',
}


def _norm(value):
    return re.sub(r'[^a-z0-9]+', '_', str(value or '').strip().lower()).strip('_')


def _to_int(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value)
    digits = re.sub(r'[^0-9]', '', str(value))
    return int(digits) if digits else None


def _cell_text(value):
    if value is None:
        return ''
    return str(value).strip()


def split_areas(value):
    text = _cell_text(value)
    if not text:
        return []
    parts = re.split(r'[,;/|]', text)
    return [part.strip() for part in parts if part.strip()]


def read_rows_from_workbook(fileobj):
    from openpyxl import load_workbook

    workbook = load_workbook(fileobj, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            return []

        col_map = {}
        for index, cell in enumerate(header):
            key = HEADER_MAP.get(_norm(cell))
            if key and key not in col_map:
                col_map[key] = index

        rows = []
        for raw in rows_iter:
            if raw is None:
                continue
            row = {}
            for key, index in col_map.items():
                if index < len(raw):
                    row[key] = raw[index]
            rows.append(row)
        return rows
    finally:
        workbook.close()


def import_products_from_rows(rows, update_existing=True):
    summary = {
        'total': 0,
        'created': 0,
        'updated': 0,
        'skipped': 0,
        'areas_created': 0,
        'errors': [],
    }

    for row_number, row in enumerate(rows, start=2):
        summary['total'] += 1
        name = _cell_text(row.get('name'))
        if not name:
            summary['skipped'] += 1
            continue

        group = _cell_text(row.get('group'))
        presentation = _cell_text(row.get('presentation'))
        sequence_number = _to_int(row.get('sequence_number'))
        sku = _cell_text(row.get('sku'))
        areas = split_areas(row.get('location'))

        try:
            with transaction.atomic():
                product = Product.objects.filter(name__iexact=name).first()
                if product is not None and not update_existing:
                    summary['skipped'] += 1
                    continue

                created = product is None
                if created:
                    product = Product(name=name)

                if group:
                    product.group = group
                if presentation:
                    product.presentation = presentation
                if sequence_number is not None:
                    product.sequence_number = sequence_number
                if sku:
                    product.sku = sku
                if created:
                    product.unit = 'unidad'
                    if sequence_number is None:
                        product.sequence_number = 10
                product.save()

                for area_name in areas:
                    category, area_created = Category.objects.get_or_create(name=area_name)
                    if area_created:
                        summary['areas_created'] += 1
                    product.location.add(category)

                if created:
                    summary['created'] += 1
                else:
                    summary['updated'] += 1
        except Exception as exc:  # noqa: BLE001
            summary['errors'].append(f'Fila {row_number} ({name}): {exc}')

    return summary


def import_products_from_workbook(fileobj, update_existing=True):
    rows = read_rows_from_workbook(fileobj)
    return import_products_from_rows(rows, update_existing=update_existing)
